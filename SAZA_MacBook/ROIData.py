#!/usr/bin/env python
# coding: utf-8

# In[13]:


def ROIData_edit(data_folder,TMIN,TMAX):

    import numpy as np
    import os, re, csv, tkinter, shutil
    import pandas as pd
    from itertools import compress
    from InROI import InROI
    from tkinter import filedialog
    from openpyxl import load_workbook
    import xlsxwriter

    print('ROIData function BEGIN ------')
        
    ## define var
    # basic time section
    tmin = [0,180,360,540,720,900,1080,1260]
    tmax = [180,360,540,720,900,1080,1260,1440]
    tmin = TMIN
    tmax = TMAX
    filename = np.array([])
    ROI_left = np.array([[0,50], [0,75], [15,50], [15,75]])
    ROI_right = np.array([[15,50],[15,75],[30,50],[30,75]])
    ROI_box = [ROI_left,ROI_right]
    #ROI_top = np.array([[0,0],[0,50],[30,0],[30,50]])
    #ROI_box = [ROI_left,ROI_right,ROI_top]
    data_f = '%s/%s/FishData' %(os.getcwd(),data_folder)
    ROI_folder = '%s/%s' %(os.getcwd(),data_folder)
    
    #define file name by time section, default ROI boxes are left and right
    for region in ['L','R']:
        for i in range(len(tmin)):
            filename = np.append(filename,['%sto%s_%s'%(str(tmin[i]),str(tmax[i]),region)])
    filename = filename.reshape(len(ROI_box),len(tmin))
    #filename = np.array(['pre_left','STIM_left','post_left','pre_right','STIM_right','post_right'])

    
    
    # Define function to sort .csv files in order
    _nsre = re.compile('([0-9]+)')
    def natural_sort_key(s):
        return [int(text) if text.isdigit() else text.lower()
                for text in re.split(_nsre, s)]

    csvfiles = []
    for File in os.listdir(data_f):
        if File.endswith('.csv'):
            csvfiles.append(os.path.join(data_f, File))
    csvfiles.sort(key=natural_sort_key)
    data = pd.DataFrame(columns = ['Fish Number', 'TIn','TOut','BoundaryCrossings',\
                                   'MeanVIn','MeanVOut','StdVIn','StdVOut','ENTRIES',\
                                   'Mean T per entry','% time in record','sanity'])

    # cal and make file
    if os.path.exists(ROI_folder + '/%s_ROI_edit.xlsx'%data_folder):
        os.remove(ROI_folder+'/%s_ROI_edit.xlsx'% data_folder)

    with pd.ExcelWriter(ROI_folder + '/%s_ROI_edit.xlsx'%data_folder,engine = 'xlsxwriter') as writer:
        for m in range(len(ROI_box)):
            # m indicates the programming ROI 
            Coordinates = ROI_box[m]
            for n in range (len(tmin)):
                TMin, TMax = tmin[n], tmax[n]
                for i in range(1, len(csvfiles)+1):
                    print('Fish..'+str(i))
                    # Extract time spent
                    File = pd.read_csv(csvfiles[i-1])
                    cnames = File.columns.tolist()
                    time = File['T']
                    index = 'Fish %s' % i

                    TMin1 = round(File['Sampling_Rate'][0]*TMin)
                    TMax1 = round(File['Sampling_Rate'][0]*TMax)
                    data_in, data_out = InROI(File, Coordinates, TMin1, TMax1)

                    TIn = 0
                    TOut = 0
                    Changes = 0
                    entry = 0
                    for j in range(1, len(data_in)):
                        if data_in[j-1] != data_in[j]:
                            Changes += 1
                    VIn=list(compress(File['V'],data_in[:-1]))
                    VOut=list(compress(File['V'],data_out[:-1]))
                    if VIn:
                        if np.isnan(VIn[0]):
                            del VIn[0]
                    if VOut:
                        if np.isnan(VOut[0]):
                            del VOut[0]
                    dT = File['dT'][1:]
                    TIn += sum(list(compress(dT,data_in[:-1])))
                    TOut += sum(list(compress(dT,data_out[:-1])))
                    sanity = TIn+TOut
                    perc_time_rec = round((TIn/sanity)*100,4)
                    entry = int(Changes/2)
                    if entry ==0:
                        Mean_t_per_entry= 'No entry'
                    else:
                        Mean_t_per_entry = round(TIn/entry,5)

                    data = pd.concat([data, pd.DataFrame({'Fish Number': index, 'TIn':TIn, 'TOut':TOut, 'BoundaryCrossings':Changes,\
                            'MeanVIn':np.mean(VIn), 'MeanVOut':np.mean(VOut),\
                            'StdVIn':np.std(VIn), 'StdVOut':np.std(VOut),'ENTRIES':entry,\
                            'Mean T per entry': Mean_t_per_entry,'% time in ROI':perc_time_rec,\
                            'sanity':sanity}, index=[0])], ignore_index=True)

                # when done each ROI, save data to specific sheet
                data.to_excel(writer, sheet_name=filename[m][n], index = False)
                worksheet = writer.sheets[filename[m][n]]
                worksheet.set_column('A:D',10)
                worksheet.set_column('D:L',13)
                data.drop(data.index, inplace = True)


# In[14]:


def ROIData_180sec(data_folder):

    import numpy as np
    import os, re, csv, tkinter, shutil
    import pandas as pd
    from itertools import compress
    from InROI import InROI
    from tkinter import filedialog
    from openpyxl import load_workbook
    import xlsxwriter

    print('ROIData function BEGIN ------')
        
    ## define var
    # basic time section
    tmin = [0,180,360,540,720,900,1080,1260]
    tmax = [180,360,540,720,900,1080,1260,1440]
    filename = np.array([])
    ROI_left = np.array([[0,50], [0,75], [15,50], [15,75]])
    ROI_right = np.array([[15,50],[15,75],[30,50],[30,75]])
    ROI_box = [ROI_left,ROI_right]
    #ROI_top = np.array([[0,0],[0,50],[30,0],[30,50]])
    #ROI_box = [ROI_left,ROI_right,ROI_top]
    data_f = '%s/%s/FishData' %(os.getcwd(),data_folder)
    ROI_folder = '%s/%s' %(os.getcwd(),data_folder)
    
    #define file name by time section, default ROI boxes are left and right
    for region in ['L','R']:
        for i in range(len(tmin)):
            filename = np.append(filename,['%sto%s_%s'%(str(tmin[i]),str(tmax[i]),region)])
    filename = filename.reshape(len(ROI_box),len(tmin))
    #filename = np.array(['pre_left','STIM_left','post_left','pre_right','STIM_right','post_right'])

    
    
    # Define function to sort .csv files in order
    _nsre = re.compile('([0-9]+)')
    def natural_sort_key(s):
        return [int(text) if text.isdigit() else text.lower()
                for text in re.split(_nsre, s)]

    csvfiles = []
    for File in os.listdir(data_f):
        if File.endswith('.csv'):
            csvfiles.append(os.path.join(data_f, File))
    csvfiles.sort(key=natural_sort_key)
    data = pd.DataFrame(columns = ['Fish Number', 'TIn','TOut','BoundaryCrossings',\
                                   'MeanVIn','MeanVOut','StdVIn','StdVOut','ENTRIES',\
                                   'Mean T per entry','% time in record','sanity'])

    # cal and make file
    if os.path.exists(ROI_folder + '/%s_ROI_180sec.xlsx'%data_folder):
        os.remove(ROI_folder+'/%s_ROI_180sec.xlsx'% data_folder)

    with pd.ExcelWriter(ROI_folder + '/%s_ROI_180sec.xlsx'%data_folder,engine = 'xlsxwriter') as writer:
        for m in range(len(ROI_box)):
            # m indicates the programming ROI 
            Coordinates = ROI_box[m]
            for n in range (len(tmin)):
                TMin, TMax = tmin[n], tmax[n]
                for i in range(1, len(csvfiles)+1):
                    print('Fish..'+str(i))
                    # Extract time spent
                    File = pd.read_csv(csvfiles[i-1])
                    cnames = File.columns.tolist()
                    time = File['T']
                    index = 'Fish %s' % i

                    TMin1 = round(File['Sampling_Rate'][0]*TMin)
                    TMax1 = round(File['Sampling_Rate'][0]*TMax)
                    data_in, data_out = InROI(File, Coordinates, TMin1, TMax1)

                    TIn = 0
                    TOut = 0
                    Changes = 0
                    entry = 0
                    for j in range(1, len(data_in)):
                        if data_in[j-1] != data_in[j]:
                            Changes += 1
                    VIn=list(compress(File['V'],data_in[:-1]))
                    VOut=list(compress(File['V'],data_out[:-1]))
                    if VIn:
                        if np.isnan(VIn[0]):
                            del VIn[0]
                    if VOut:
                        if np.isnan(VOut[0]):
                            del VOut[0]
                    dT = File['dT'][1:]
                    TIn += sum(list(compress(dT,data_in[:-1])))
                    TOut += sum(list(compress(dT,data_out[:-1])))
                    sanity = TIn+TOut
                    perc_time_rec = round((TIn/sanity)*100,4)
                    entry = int(Changes/2)
                    if entry ==0:
                        Mean_t_per_entry= 'No entry'
                    else:
                        Mean_t_per_entry = round(TIn/entry,5)

                    data = pd.concat([data, pd.DataFrame({'Fish Number': index, 'TIn':TIn, 'TOut':TOut, 'BoundaryCrossings':Changes,\
                            'MeanVIn':np.mean(VIn), 'MeanVOut':np.mean(VOut),\
                            'StdVIn':np.std(VIn), 'StdVOut':np.std(VOut),'ENTRIES':entry,\
                            'Mean T per entry': Mean_t_per_entry,'% time in ROI':perc_time_rec,\
                            'sanity':sanity}, index=[0])], ignore_index=True)

                # when done each ROI, save data to specific sheet
                data.to_excel(writer, sheet_name=filename[m][n], index = False)
                worksheet = writer.sheets[filename[m][n]]
                worksheet.set_column('A:D',10)
                worksheet.set_column('D:L',13)
                data.drop(data.index, inplace = True)


# In[15]:


def ROIData_std(data_folder):

    import numpy as np
    import os, re, csv, tkinter, shutil
    import pandas as pd
    from itertools import compress
    from InROI import InROI
    from tkinter import filedialog
    from openpyxl import load_workbook
    import xlsxwriter

    print('ROIData function BEGIN ------')
        
    ## define var
    # basic time section
    tmin = [0,180,720,1260]
    tmax = [180,720,1260,1440]
    filename = np.array([])
    ROI_left = np.array([[0,50], [0,75], [15,50], [15,75]])
    ROI_right = np.array([[15,50],[15,75],[30,50],[30,75]])
    ROI_box = [ROI_left,ROI_right]
    #ROI_top = np.array([[0,0],[0,50],[30,0],[30,50]])
    #ROI_box = [ROI_left,ROI_right,ROI_top]
    data_f = '%s/%s/FishData' %(os.getcwd(),data_folder)
    ROI_folder = '%s/%s' %(os.getcwd(),data_folder)
    
    #define file name by time section, default ROI boxes are left and right
    for region in ['L','R']:
        for i in range(len(tmin)):
            filename = np.append(filename,['%sto%s_%s'%(str(tmin[i]),str(tmax[i]),region)])
    filename = filename.reshape(len(ROI_box),len(tmin))
    #filename = np.array(['pre_left','STIM_left','post_left','pre_right','STIM_right','post_right'])

    
    
    # Define function to sort .csv files in order
    _nsre = re.compile('([0-9]+)')
    def natural_sort_key(s):
        return [int(text) if text.isdigit() else text.lower()
                for text in re.split(_nsre, s)]

    csvfiles = []
    for File in os.listdir(data_f):
        if File.endswith('.csv'):
            csvfiles.append(os.path.join(data_f, File))
    csvfiles.sort(key=natural_sort_key)
    data = pd.DataFrame(columns = ['Fish Number', 'TIn','TOut','BoundaryCrossings',\
                                   'MeanVIn','MeanVOut','StdVIn','StdVOut','ENTRIES',\
                                   'Mean T per entry','% time in record','sanity'])

    # cal and make file
    if os.path.exists(ROI_folder + '/%s_ROI_std.xlsx'%data_folder):
        os.remove(ROI_folder+'/%s_ROI_std.xlsx'% data_folder)

    with pd.ExcelWriter(ROI_folder + '/%s_ROI_std.xlsx'%data_folder,engine = 'xlsxwriter') as writer:
        for m in range(len(ROI_box)):
            # m indicates the programming ROI 
            Coordinates = ROI_box[m]
            for n in range (len(tmin)):
                TMin, TMax = tmin[n], tmax[n]
                for i in range(1, len(csvfiles)+1):
                    print('Fish..'+str(i))
                    # Extract time spent
                    File = pd.read_csv(csvfiles[i-1])
                    cnames = File.columns.tolist()
                    time = File['T']
                    index = 'Fish %s' % i

                    TMin1 = round(File['Sampling_Rate'][0]*TMin)
                    TMax1 = round(File['Sampling_Rate'][0]*TMax)
                    data_in, data_out = InROI(File, Coordinates, TMin1, TMax1)

                    TIn = 0
                    TOut = 0
                    Changes = 0
                    entry = 0
                    for j in range(1, len(data_in)):
                        if data_in[j-1] != data_in[j]:
                            Changes += 1
                    VIn=list(compress(File['V'],data_in[:-1]))
                    VOut=list(compress(File['V'],data_out[:-1]))
                    if VIn:
                        if np.isnan(VIn[0]):
                            del VIn[0]
                    if VOut:
                        if np.isnan(VOut[0]):
                            del VOut[0]
                    dT = File['dT'][1:]
                    TIn += sum(list(compress(dT,data_in[:-1])))
                    TOut += sum(list(compress(dT,data_out[:-1])))
                    sanity = TIn+TOut
                    perc_time_rec = round((TIn/sanity)*100,4)
                    entry = int(Changes/2)
                    if entry ==0:
                        Mean_t_per_entry= 'No entry'
                    else:
                        Mean_t_per_entry = round(TIn/entry,5)

                    data = pd.concat([data, pd.DataFrame({'Fish Number': index, 'TIn':TIn, 'TOut':TOut, 'BoundaryCrossings':Changes,\
                            'MeanVIn':np.mean(VIn), 'MeanVOut':np.mean(VOut),\
                            'StdVIn':np.std(VIn), 'StdVOut':np.std(VOut),'ENTRIES':entry,\
                            'Mean T per entry': Mean_t_per_entry,'% time in ROI':perc_time_rec,\
                            'sanity':sanity}, index=[0])], ignore_index=True)

                # when done each ROI, save data to specific sheet
                data.to_excel(writer, sheet_name=filename[m][n], index = False)
                worksheet = writer.sheets[filename[m][n]]
                worksheet.set_column('A:D',10)
                worksheet.set_column('D:L',13)
                data.drop(data.index, inplace = True)


# In[16]:


def ROIData_basic(data_folder):

    import numpy as np
    import os, re, csv, tkinter, shutil
    import pandas as pd
    from itertools import compress
    from InROI import InROI
    from tkinter import filedialog
    from openpyxl import load_workbook
    import xlsxwriter

    print('ROIData function BEGIN ------')
        
    ## define var
    tmin = [0,180,1260]
    tmax = [180,1260,1440]
    #filename = np.array(['pre_left','STIM_left','post_left','pre_right','STIM_right','post_right'])
    ROI_left = np.array([[0,50], [0,75], [15,50], [15,75]])
    ROI_right = np.array([[15,50],[15,75],[30,50],[30,75]])
    ROI_box = [ROI_left,ROI_right]
    #ROI_top = np.array([[0,0],[0,50],[30,0],[30,50]])
    #ROI_box = [ROI_left,ROI_right,ROI_top]
    pre_path = os.getcwd()
    data_f = '%s/%s/FishData' %(pre_path,data_folder)
    ROI_folder = '%s/%s' %(pre_path,data_folder)
    
    filename = np.array([])
    for region in ['L','R']:
        for i in range(len(tmin)):
            filename = np.append(filename,['%sto%s_%s'%(str(tmin[i]),str(tmax[i]),region)])
    filename = filename.reshape(len(ROI_box),len(tmin))
    
    # Define function to sort .csv files in order
    _nsre = re.compile('([0-9]+)')
    def natural_sort_key(s):
        return [int(text) if text.isdigit() else text.lower()
                for text in re.split(_nsre, s)]

    csvfiles = []
    for File in os.listdir(data_f):
        if File.endswith('.csv'):
            csvfiles.append(os.path.join(data_f, File))
    csvfiles.sort(key=natural_sort_key)
    data = pd.DataFrame(columns = ['Fish Number', 'TIn','TOut','BoundaryCrossings',\
                                   'MeanVIn','MeanVOut','StdVIn','StdVOut','ENTRIES',\
                                   'Mean T per entry','% time in record','sanity'])

    # cal and make file
    if os.path.exists(ROI_folder + '/%s_ROI_basic.xlsx'%data_folder):
        os.remove(ROI_folder+'/%s_ROI_basic.xlsx'% data_folder)

    with pd.ExcelWriter(ROI_folder + '/%s_ROI_basic.xlsx'%data_folder,engine = 'xlsxwriter') as writer:
        for m in range(len(ROI_box)):
            # m indicates the programming ROI 
            Coordinates = ROI_box[m]
            for n in range (len(tmin)):
                TMin, TMax = tmin[n], tmax[n]
                for i in range(1, len(csvfiles)+1):
                    print('Fish..'+str(i))
                    # Extract time spent
                    File = pd.read_csv(csvfiles[i-1])
                    cnames = File.columns.tolist()
                    time = File['T']
                    index = 'Fish %s' % i

                    TMin1 = round(File['Sampling_Rate'][0]*TMin)
                    TMax1 = round(File['Sampling_Rate'][0]*TMax)
                    data_in, data_out = InROI(File, Coordinates, TMin1, TMax1)

                    TIn = 0
                    TOut = 0
                    Changes = 0
                    entry = 0
                    for j in range(1, len(data_in)):
                        if data_in[j-1] != data_in[j]:
                            Changes += 1
                    VIn=list(compress(File['V'],data_in[:-1]))
                    VOut=list(compress(File['V'],data_out[:-1]))
                    if VIn:
                        if np.isnan(VIn[0]):
                            del VIn[0]
                    if VOut:
                        if np.isnan(VOut[0]):
                            del VOut[0]
                    dT = File['dT'][1:]
                    TIn += sum(list(compress(dT,data_in[:-1])))
                    TOut += sum(list(compress(dT,data_out[:-1])))
                    sanity = TIn+TOut
                    perc_time_rec = round((TIn/sanity)*100,4)
                    entry = int(Changes/2)
                    if entry ==0:
                        Mean_t_per_entry= 'No entry'
                    else:
                        Mean_t_per_entry = round(TIn/entry,5)

                    data = pd.concat([data, pd.DataFrame({'Fish Number': index, 'TIn':TIn, 'TOut':TOut, 'BoundaryCrossings':Changes,\
                            'MeanVIn':np.mean(VIn), 'MeanVOut':np.mean(VOut),\
                            'StdVIn':np.std(VIn), 'StdVOut':np.std(VOut),'ENTRIES':entry,\
                            'Mean T per entry': Mean_t_per_entry,'% time in ROI':perc_time_rec,\
                            'sanity':sanity}, index=[0])], ignore_index=True)

                # when done each ROI, save data to specific sheet
                data.to_excel(writer, sheet_name=filename[m][n], index = False)
                worksheet = writer.sheets[filename[m][n]]
                worksheet.set_column('A:D',10)
                worksheet.set_column('D:L',13)
                data.drop(data.index, inplace = True)

